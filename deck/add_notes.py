"""
deck/add_notes.py
=================
Attach presenter-view speaker notes to the built deck.

Run AFTER `node deck/main12.js`. pptxgenjs writes the slides; this adds the notes,
because a rebuild would otherwise wipe anything added by hand.

    node deck/main12.js && python deck/add_notes.py

These are CUES, not a script - short enough to actually read while talking. The full
narration is in PRESENTATION_GUIDE.md.
"""
from pptx import Presentation

NOTES = {
 1: "30s. Do NOT give the conclusion yet - slides 2-5 set up the problem.",
 2: "90s. Three beats: (1) what glucose is and why both directions are dangerous, "
    "(2) today it means breaking skin, 590M adults, (3) the amber box - the biology "
    "supports DETECTING an event, not reading the exact number. Classification vs "
    "regression. This project attempted the number.",
 3: "90s. IEEE TNNLS, a serious journal. 1.49 mmol/L approaches approved CGM accuracy. "
    "They published code but NOT data - so nobody had independently checked it. "
    "That is what a reproduction is for. Walk the three stage boxes left to right.",
 4: "90s. SLOW DOWN on the dark box. 2.519 is the standard deviation of glucose, and "
    "therefore the error of a model that always guesses the average. It is the score "
    "for learning NOTHING. Most papers never report it; we report it beside every "
    "result. Everything after this slide depends on them having this.",
 5: "2 min, densest slide. START WITH THE ANALOGY at bottom left - three doctors, two "
    "trained together so averaging double-counts, one is more experienced. A weighted "
    "average handles the second but NOT the first. THEN the formulas. Note the weights "
    "(.300,.342,.358) are not the densities (.30,.25,.20) - that gap is the integral "
    "doing something a weighted average cannot.",
 6: "60s. CUT THIS FIRST if short on time. Key line: a heartbeat mixes fast (QRS, 0.1s) "
    "and slow (T wave, 0.3s) things, so measuring the whole window blends them; the "
    "wavelet separates them into 8 layers. Then 8 x 10 = 80 per signal. Mention QTc "
    "median 0.405s vs textbook 0.35-0.44 as validation.",
 7: "90s. THE most important concept in the talk. Walk the three boxes. Then the key "
    "insight, said slowly: consecutive windows from one person are very similar, so a "
    "random split lets the model recognise the PERSON and recall their typical glucose. "
    "That is not learning about glucose - but it scores beautifully. Close with the "
    "2026 study: 0.60 under random, -0.08 under participant-aware.",
 8: "90s. DEFINE R-SQUARED FIRST - 1 perfect, 0 = same as guessing the average, "
    "NEGATIVE = worse than guessing. Then: +0.149 random, -0.223 subject-aware, "
    "-0.097 leave-one-out. Same data, same code, only the split changed. "
    "RAISE THE ORDERING YOURSELF: LOSO is better than subject-aware because it trains "
    "on 9 of 10 rather than 4 of 5. Saying it first is far stronger than being caught.",
 9: "90s. YOUR BEST SLIDE. Trace the diagonal with your hand - that is what a working "
    "model looks like. Then trace the horizontal band - whether the truth was 4 or 20, "
    "it predicts 6-9. That is what 'no better than the average' looks like drawn. "
    "THEN the sting: that same flat band scores 89.4% Zone A+B. PAUSE. Let it land.",
10: "75s. Point at the LAST row. The baseline predicts a constant, explains zero "
    "variance, and has BOTH the best R-squared (-0.074) and the best clinical score "
    "(90.0%). Why: most readings sit in a narrow band, so a constant already lands in "
    "A and B. The grid catches DANGEROUS errors, not USELESS ones. Rule: never report "
    "a zone score without a baseline beside it.",
11: "45s, quick and visual. Red bar = the constant predictor. LONGEST on the left "
    "(best clinical score), SHORTEST on the right (least-bad R-squared). One metric "
    "says it wins, the other says nothing works. If you only report the left chart, "
    "you would ship this.",
12: "90s. Start with the two clues: weighted average = plain average byte-identical "
    "(impossible if weights differ), and every density at exactly the 0.01 floor. "
    "Then the theorem: equal densities make the measure symmetric, so the integral "
    "becomes a fixed function of RANK. Weak models -> lambda > 0 -> beta > 1 -> weight "
    "grows toward the SMALLEST. At g=0.01 that is 89.5% on the minimum. "
    "Confirmation: corr 0.998 with min(). DO NOT SKIP the dark box - it is not only a "
    "bug, it fires exactly when the models cannot generalise.",
13: "60s. Two points. LEFT: every bar is past the dashed baseline - no subset of the "
    "193 features recovers glucose. And the ordering runs the WRONG WAY: fewer features "
    "score better, which is what fitting noise looks like. RIGHT: the worst models by "
    "R-squared carry the HIGHEST clinical score. Not weak - inverted. A model predicting "
    "a narrow band near the average is never dangerous, just useless - and the grid only "
    "catches danger.",
14: "60s. CUT IF SHORT - the result is already on slide 12. Three substitutions: "
    "apply Sugeno j times, subtract consecutive terms, check the geometric series sums "
    "to 1. The red box is the point: weak models force beta > 1, which piles weight "
    "on the smallest prediction.",
15: "60s. RANSAC = random sample consensus, unseeded, so identical input gives "
    "different output - 5 distinct values in 30 calls. Up to 72% of correlation-"
    "dimension values changed. POINT AT THE SECOND NUMBER: 177 other features "
    "bit-identical - that is the control proving it was the slope fit, not anything "
    "else we changed. Found only because we tested determinism, which almost nobody does.",
16: "90s. Four contributions, then the amber box. Say the limits CONFIDENTLY, not "
    "apologetically: cannot claim the method never works, only that it does not here "
    "under honest evaluation and that its fusion was inactive. Stating limitations is "
    "the point of a reproduction, not a weakness.",
17: "15s. 'Full write-up with every derivation is in the repository.' Then STOP "
    "TALKING. Leave this up, or go back to slide 15 for questions.",
}

DECK = "BTP_presentation.pptx"

if __name__ == "__main__":
    p = Presentation(DECK)
    for i, s in enumerate(p.slides, 1):
        if i in NOTES:
            s.notes_slide.notes_text_frame.text = NOTES[i]
    p.save(DECK)

    check = Presentation(DECK)
    missing = [i for i, s in enumerate(check.slides, 1)
               if not (s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())]
    print(f"{len(check.slides)} slides; without notes: {missing or 'none'}")
